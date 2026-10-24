"""
Document text/table extraction worker for the LR billing platform.

Usage:  python -I extract.py <file-path>
Prints one JSON object to stdout:
  {"method": str, "pages": [{"page": int, "text": str, "tables": [[[cell,...],...],...]}], "warnings": [str]}

Document contents are treated strictly as data: macros are never executed,
formulas are read as cached values only, and nothing in a file is evaluated.
"""
import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore", category=FutureWarning)

MAX_PAGES = 400


def undouble(s):
    """Some PDFs draw bold text twice, so table extraction yields 'CCoonnttrraacctt'.
    Collapse words whose characters are all doubled; leave everything else untouched."""
    if not s:
        return s

    def fix(m):
        w = m.group(0)
        # Only words with letters: amounts such as "2200" or "3300" must never change.
        if len(w) >= 4 and len(w) % 2 == 0 and re.search("[A-Za-z]", w) and w[0::2] == w[1::2]:
            return w[0::2]
        return w

    return re.sub(r"\S+", fix, s)


def pdf_extract(path):
    warnings = []
    pages = []
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    plumber = pdfplumber.open(path) if pdfplumber else None
    ocr_needed = []
    for i, page in enumerate(doc):
        if i >= MAX_PAGES:
            warnings.append(f"Only the first {MAX_PAGES} pages were processed")
            break
        text = page.get_text("text") or ""
        tables = []
        if plumber is not None:
            try:
                for t in plumber.pages[i].extract_tables():
                    tables.append([[("" if c is None else undouble(str(c).strip())) for c in row] for row in t])
            except Exception as e:  # table detection is best-effort
                warnings.append(f"Page {i + 1}: table extraction failed ({e.__class__.__name__})")
        if len(text.strip()) < 20:
            ocr_needed.append(i)
            try:
                tp = page.get_textpage_ocr(language="eng", dpi=300, full=True)
                text = page.get_text("text", textpage=tp)
                warnings.append(f"Page {i + 1}: text recognised by OCR - verify values carefully")
            except Exception:
                warnings.append(f"Page {i + 1}: no text layer and OCR (Tesseract) is not available - page is probably scanned")
        pages.append({"page": i + 1, "text": text, "tables": tables})
    if plumber is not None:
        plumber.close()
    return {"method": "pymupdf+pdfplumber", "pages": pages, "warnings": warnings, "pageCount": len(doc)}


def xlsx_extract(path):
    import openpyxl

    # data_only=True reads cached formula results; keep_vba is never enabled.
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pages = []
    for n, ws in enumerate(wb.worksheets):
        rows = []
        lines = []
        for r in ws.iter_rows():
            vals = []
            for c in r:
                v = c.value
                if v is None:
                    vals.append("")
                elif hasattr(v, "isoformat"):
                    vals.append(v.isoformat())
                else:
                    vals.append(str(v))
            if any(vals):
                rows.append(vals)
                cells = [f"{c.coordinate}={vals[j]}" for j, c in enumerate(r) if vals[j] and hasattr(c, "coordinate")]
                lines.append(" | ".join(cells) if cells else " | ".join(vals))
        pages.append({"page": n + 1, "sheet": ws.title, "text": "\n".join(lines), "tables": [rows]})
    return {"method": "openpyxl", "pages": pages, "warnings": [], "pageCount": len(pages)}


def xls_extract(path):
    """Legacy Excel 97-2003 workbooks (cell values only; macros are never run)."""
    import xlrd

    book = xlrd.open_workbook(path)
    pages = []
    for n, sh in enumerate(book.sheets()):
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    y, mo, d, h, mi, _ = xlrd.xldate_as_tuple(cell.value, book.datemode)
                    vals.append(f"{y:04d}-{mo:02d}-{d:02d}" + (f" {h:02d}:{mi:02d}" if (h or mi) else ""))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    v = cell.value
                    vals.append(str(int(v)) if float(v).is_integer() else str(v))
                else:
                    vals.append(str(cell.value).strip())
            rows.append(vals)
        text = "\n".join(" | ".join(r) for r in rows if any(r))
        pages.append({"page": n + 1, "sheet": sh.name, "text": text, "tables": [rows]})
    return {"method": "xlrd", "pages": pages, "warnings": [], "pageCount": len(pages)}


def docx_extract(path):
    import docx

    d = docx.Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    for t in tables:
        text += "\n" + "\n".join(" | ".join(r) for r in t)
    return {"method": "python-docx", "pages": [{"page": 1, "text": text, "tables": tables}], "warnings": ["Word documents have no page numbers; sources are shown as page 1"], "pageCount": 1}


def image_extract(path):
    import fitz

    try:
        doc = fitz.open(path)
        pdfbytes = doc.convert_to_pdf()
        pdf = fitz.open("pdf", pdfbytes)
        page = pdf[0]
        tp = page.get_textpage_ocr(language="eng", dpi=300, full=True)
        text = page.get_text("text", textpage=tp)
        return {"method": "ocr", "pages": [{"page": 1, "text": text, "tables": []}], "warnings": ["Text recognised by OCR - verify values carefully"], "pageCount": 1}
    except Exception:
        return {"method": "none", "pages": [{"page": 1, "text": "", "tables": []}], "warnings": ["OCR (Tesseract) is not available; enter values from this image manually"], "pageCount": 1}


def main():
    if len(sys.argv) != 2:
        print(json.dumps({"error": "usage: extract.py <file>"}))
        return 2
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            out = pdf_extract(path)
        elif ext in (".xlsx", ".xlsm"):
            out = xlsx_extract(path)
        elif ext == ".xls":
            out = xls_extract(path)
        elif ext == ".docx":
            out = docx_extract(path)
        elif ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            out = image_extract(path)
        elif ext == ".csv":
            with open(path, encoding="utf-8", errors="replace") as f:
                text = f.read()
            out = {"method": "text", "pages": [{"page": 1, "text": text, "tables": []}], "warnings": [], "pageCount": 1}
        else:
            out = {"error": f"Unsupported file type {ext}"}
    except Exception as e:
        out = {"error": f"{e.__class__.__name__}: {e}"}
    # Always UTF-8, whatever the console code page (e.g. cp1252 on Windows).
    sys.stdout.buffer.write(json.dumps(out, ensure_ascii=False).encode("utf-8"))
    sys.stdout.flush()
    return 0 if "error" not in out else 1


if __name__ == "__main__":
    sys.exit(main())
